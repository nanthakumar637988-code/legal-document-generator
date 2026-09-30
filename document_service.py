from io import BytesIO
from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches,Pt
from fpdf import FPDF
from utils.text_utils import sanitize_text
BASE_DIR=Path(__file__).resolve().parents[1]; LOGO_PATH=BASE_DIR/"assets"/"legal_ease_logo.png"
def format_txt(text): return sanitize_text(text).encode("utf-8")
def format_docx(text,doc_type):
    d=Document(); s=d.sections[0]; s.top_margin=Inches(.7); s.bottom_margin=Inches(.7); s.left_margin=Inches(.8); s.right_margin=Inches(.8)
    if LOGO_PATH.exists():
        p=s.header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(LOGO_PATH),width=Inches(1.25))
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(sanitize_text(doc_type).upper()); r.bold=True; r.font.name="Times New Roman"; r.font.size=Pt(16)
    for line in sanitize_text(text).splitlines():
        if not line.strip(): d.add_paragraph(); continue
        p=d.add_paragraph(); p.paragraph_format.space_after=Pt(6); r=p.add_run(line.strip()); r.font.name="Times New Roman"; r.font.size=Pt(12 if re.match(r"^\d+\.\s+[A-Z][A-Z0-9 /&-]+$",line.strip()) or line.strip() in {"NOTICE","PARTIES","SIGNATURES"} else 11); r.bold=(r.font.size.pt==12)
    f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=f.add_run("LegalEase - AI-generated draft. Review before use."); rr.font.size=Pt(8)
    out=BytesIO(); d.save(out); return out.getvalue()
class LegalEasePDF(FPDF):
    def __init__(self,doc_type): super().__init__(); self.doc_type=doc_type; self.set_auto_page_break(True,18)
    def header(self):
        if LOGO_PATH.exists():
            try: self.image(str(LOGO_PATH),x=92,y=8,w=26); self.ln(17)
            except Exception: self.ln(4)
        self.set_font("Helvetica","B",11); self.cell(0,7,self.doc_type.upper(),align="C"); self.ln(9)
    def footer(self): self.set_y(-14); self.set_font("Helvetica","",8); self.cell(0,8,"LegalEase - AI-generated draft. Review before use.",align="C")
def format_pdf(text,doc_type):
    pdf=LegalEasePDF(doc_type); pdf.add_page(); pdf.set_font("Helvetica","",10)
    for line in sanitize_text(text).splitlines():
        if not line.strip(): pdf.ln(3); continue
        heading=bool(re.match(r"^\d+\.\s+[A-Z][A-Z0-9 /&-]+$",line.strip()) or line.strip() in {"NOTICE","PARTIES","SIGNATURES"})
        pdf.set_font("Helvetica","B" if heading else "",11 if heading else 10); pdf.multi_cell(pdf.epw,6 if heading else 5.5,line.strip(),new_x="LMARGIN",new_y="NEXT")
    return bytes(pdf.output())
def export_filename(doc_type,ext):
    name=re.sub(r"[^A-Za-z0-9_-]+","_",doc_type.strip()).strip("_") or "legal_document"
    return f"{name}.{ext}"
